#!/usr/bin/env python3
"""Build a Word report from saved iScholar roadshow rehearsal evidence."""

import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "doc"
EVIDENCE = DOC / "roadshow-evidence"
DATA = json.loads((EVIDENCE / "rehearsal-data.json").read_text(encoding="utf-8"))
OUTPUT = DOC / "roadshow-rehearsal-report.docx"

NAVY, TEAL, PALE, PALE_BLUE = "102B45", "0F8B8D", "EAF4F5", "EEF3F8"
GRAY, WHITE = "536579", "FFFFFF"


def set_font(run, size=9.5, bold=False, color="24364B"):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def margins(cell, top=100, side=130, bottom=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("start", side), ("bottom", bottom), ("end", side)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def section_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), {1: 16, 2: 12}[level], True, NAVY if level == 1 else TEAL)
    return p


def body(doc, text, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text), size)
    return p


def bullet(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), 9.2)
    return p


def make_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, title in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Cm(widths[idx])
        shade(cell, NAVY)
        margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(title)), 8.5, True, WHITE)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = Cm(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if row_idx % 2 == 0:
                shade(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_font(p.add_run(str(value)), 8.2)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(1)
    return table


def callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, PALE)
    margins(cell, 150, 190, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(title), 10, True, TEAL)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    set_font(p2.add_run(text), 9.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def local_time(iso):
    dt = datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(ZoneInfo("Asia/Shanghai"))
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def elapsed(iso_start, iso_end):
    start = datetime.fromisoformat(iso_start.replace("Z", "+00:00"))
    end = datetime.fromisoformat(iso_end.replace("Z", "+00:00"))
    seconds = (end - start).total_seconds()
    minutes, seconds_rounded = divmod(round(seconds), 60)
    return f"{minutes} 分 {seconds_rounded} 秒（约 {seconds / 60:.1f} 分钟）"


def add_page_number(p):
    run = p.add_run()
    begin, instr, end = (OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar"))
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_font(run, 8, color=GRAY)


def add_screenshot_appendix(doc, shots, heading):
    doc.add_page_break()
    section_heading(doc, heading)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx in range(0, len(shots), 2):
        cells = table.add_row().cells
        for col in range(2):
            cell = cells[col]
            cell.width = Cm(8.15)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            margins(cell, 120, 100, 120)
            shot_idx = idx + col
            if shot_idx >= len(shots):
                continue
            filename, caption = shots[shot_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(EVIDENCE / filename), width=Cm(7.75))
            cp = cell.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(2)
            set_font(cp.add_run(caption), 8, True, GRAY)


# Derived facts are read only from the saved run payload and the roadshow rehearsal notes.
run = DATA["run"]
created, completed = run["createdAt"], run["completedAt"]
duration = elapsed(created, completed)
artifacts = run["artifacts"]
literature = next(a for a in artifacts if a["kind"] == "literature_brief")
design = next(a for a in artifacts if a["kind"] == "research_design")
paper_count = len((literature.get("content", {}).get("structured") or {}).get("papers") or [])
hypothesis_count = len((design.get("content", {}).get("structured") or {}).get("hypotheses") or [])
literature_text = literature.get("content", {}).get("text", "")
text_claims_45 = "共纳入 45 项" in literature_text
source_count = len(artifacts[0].get("content", {}).get("sources") or [])
approved_count = sum(a["status"] == "approved" for a in artifacts)

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
for name in ("Normal", "List Bullet", "List Number"):
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(9.5)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("iScholar  /  系统预演报告"), 8, color=GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("内部演示记录  ·  2026-10-24  ·  第 "), 8, color=GRAY)
add_page_number(footer)
set_font(footer.add_run(" 页"), 8, color=GRAY)

# Cover and metadata.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(3)
set_font(p.add_run("系统预演报告  /  2026.10"), 9, True, TEAL)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("iScholar 多智能体研究计划\n流程演示报告"), 25, True, NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
set_font(p.add_run("真实环境端到端预演 · 将运行状态与学术证据核验分开记录"), 10.5, color=GRAY)

make_table(doc, ["项目", "记录"], [
    ("预演日期", "2026 年 10 月 24 日（Asia/Shanghai）"),
    ("演示项目", DATA["projectName"]),
    ("流程运行 ID", run["id"]),
    ("运行状态", f"{run['status']}；{approved_count}/{len(artifacts)} 条产物记录为 approved"),
    ("运行时间", f"{local_time(created)} 至 {local_time(completed)}；{duration}"),
], [3.4, 13.0])

section_heading(doc, "一、执行摘要")
body(doc, "本次预演通过真实浏览器与已连接的后端服务完成项目创建、研究目标提交、多智能体协作、两个人工审批节点和最终方案保存。后台运行终态为 succeeded，共保存 5 条 approved 状态的产物记录。")
callout(doc, "结论先行", f"系统流程走到成功终态，不代表引文或研究结论已核验。一次运行从创建到完成约 {duration}，超过 10 分钟路演预留的演示时段。文献简报明确提示实时检索预算超限并要求核验引用；建议台上展示预先完成的运行记录。")

section_heading(doc, "二、预演范围与输入")
make_table(doc, ["字段", "本次记录"], [
    ("运行环境", "本机 Next.js 前端；Docker backend、worker、PostgreSQL、Redis"),
    ("模型与检索", "DeepSeek；Crossref（有成功响应，也出现限流与工具预算限制）"),
    ("研究目标", DATA["input"]["researchGoal"]),
    ("学科", DATA["input"]["discipline"]),
    ("关键词", DATA["input"]["keywords"]),
], [3.4, 13.0])
body(doc, "实际操作为：登录演示账号 → 创建项目 → 填写研究目标、学科和关键词 → 确认外部 AI 调用 → 选择候选方向 → 生成文献综述和研究设计 → 质量检查 → 审批最终方案。报告不包含账号密码。")

section_heading(doc, "三、运行结果")
lit_observation = f"结构化字段含 {paper_count} 条论文；正文称纳入 45 项，需核对数量。" if text_claims_45 else f"结构化字段含 {paper_count} 条论文；需核对正文统计口径。"
make_table(doc, ["阶段", "产物记录", "观察"], [
    ("选题建议", "topic_proposal ×2", f"两份候选主题；首份附 {source_count} 条来源记录，重试后出现同类型第二份产物。"),
    ("文献综述", "literature_brief ×1", lit_observation),
    ("研究设计", "research_design ×1", f"包含 {hypothesis_count} 条假设；本次产物正文为英文。"),
    ("最终方案", "research_plan ×1", "中文汇总研究问题、证据缺口、方法与待补材料；提示原目标与选定方向的衔接需要说明。"),
], [2.7, 3.0, 10.7])
body(doc, "最终方案把原目标从“生成式 AI 辅助反馈对修订质量的影响”扩展到反馈丰裕、评价性判断、采纳行为和无 AI 迁移。方案指出需说明目标与选题之间的关系；是否采用这一方向仍须由研究者决定。")

section_heading(doc, "四、关键节点记录")
for item in [
    "项目建立：新建演示项目并进入项目工作区。",
    "输入与授权：填写教育技术研究目标、学科与关键词，并确认外部 AI 调用。",
    "选题确认：生成 3 个候选方向，由用户选择后继续。",
    "文献与设计：生成主题综述、结构化文献条目和研究假设；运行同时记录检索限制。",
    "质量检查与终审：审查证据、假设、变量和可行性，用户批准最终方案。",
    "保存结果：后台状态 succeeded，5 条产物均显示 approved。",
]:
    bullet(doc, item, numbered=True)

section_heading(doc, "五、问题、限制与修复")
issues = [
    ("文献证据边界", "历史预演的 Crossref 部分请求返回 HTTP 429，简报提示检索预算超限且依据模型既有知识；当前简报会自动声明非已完成实时系统综述，并要求核验引用和 DOI。"),
    ("条目数量冲突", f"历史简报结构字段记录 {paper_count} 项，正文称纳入 45 项。当前输出分别标注检索来源数和结构化生成条目数，并声明正文中的其他纳入数未经核实。"),
    ("重试产物", "历史运行有两份 topic_proposal。当前按运行与逻辑步骤复用同一产物记录；内容变化写入哈希审计事件，避免重试产物一并获批。"),
    ("语言一致性", "历史运行中文输入曾产生英文阶段产物。当前工作流对 auto 语言偏好按请求文本识别中文或英文，并将语言指令用于各阶段、质量检查和汇总。"),
    ("现场时长", f"历史实时运行用时 {duration}，包含人工确认和外部请求等待。新增快速演示命令可打开已完成运行，不触发新的模型或文献检索调用。"),
    ("已修复的界面阻塞", "预演发现协作页缺少中英文按钮/AI 提示文案、首个审批后忙碌状态未释放；已补文案与状态恢复，并支持用 runId 恢复持久化运行页面。"),
]
make_table(doc, ["事项", "证据与处理"], issues, [3.1, 13.3])

section_heading(doc, "六、路演前行动项")
for item in [
    "使用 `pnpm roadshow:demo -- <已完成运行ID> <项目ID>` 展示已保存产物；不依赖现场模型与检索等待。",
    "人工核查最终方案引用、DOI、样本数和方法；确认正文与结构字段数量一致。",
    "新建一轮真实预演，复核文献声明、输出语言、产物幂等和最终审批状态。",
    "演示环境使用合成数据；调用外部模型前说明授权范围，不展示密钥或账号凭据。",
    "讲解时明确区分“运行成功/产物已保存”和“研究证据已验证”。",
]:
    bullet(doc, item)

section_heading(doc, "七、证据文件")
body(doc, "完整运行输入、状态、文本产物、结构化结果和时间戳保存在 doc/roadshow-evidence/rehearsal-data.json。以下附录收录关键页面截图；审批截图记录等待状态，页面操作控件位于工作区下方。")

shots = [
    ("00-project-created.png", "项目建立"),
    ("01-research-input.png", "研究输入就绪"),
    ("02-topic-approval.png", "选题阶段等待人工确认"),
    ("03-final-approval.png", "最终方案等待审批"),
    ("04-workflow-complete.png", "最终方案已批准"),
]
add_screenshot_appendix(doc, shots[:4], "附录 A｜关键页面截图")
add_screenshot_appendix(doc, shots[4:], "附录 A｜关键页面截图（续）")

doc.core_properties.title = "iScholar 多智能体研究计划流程演示报告"
doc.core_properties.subject = "2026-10-24 系统预演结果与路演建议"
doc.core_properties.author = "iScholar"
doc.save(OUTPUT)
print(f"Created {OUTPUT}")
